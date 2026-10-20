"""Build the plain-language Word guide from its Markdown source.

    python scripts/build_guide.py

Reads `docs/guide/GUIDE.md` and writes `Content_Processor_Guide.docx` at the
project root. The Markdown is the source of truth: edit it, then rebuild.
Supported subset: `#` title, `##`/`###` headings, paragraphs, `- ` bullets,
`1. ` numbered items, pipe tables (header row plus separator), `**bold**`.
Needs python-docx (installed with the project's `office` extra).
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "guide" / "GUIDE.md"
OUTPUT = ROOT / "Content_Processor_Guide.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1F3A5F"
BAND_FILL = "EEF2F7"
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_runs(paragraph: Paragraph, text: str) -> None:
    """Add text, honouring **bold** spans."""
    pos = 0
    for match in _BOLD.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _shade(cell_element: object, fill: str) -> None:
    tc_pr = cell_element.get_or_add_tcPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _repeat_header(row: object) -> None:
    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]  # noqa: SLF001
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def _table(document: DocxDocument, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell._tc, HEADER_FILL)  # noqa: SLF001
    _repeat_header(table.rows[0])
    for n, values in enumerate(body):
        cells = table.add_row().cells
        for i, text in enumerate(values[: len(header)]):
            cells[i].text = ""
            _add_runs(cells[i].paragraphs[0], text)
            if n % 2 == 1:
                _shade(cells[i]._tc, BAND_FILL)  # noqa: SLF001
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)


def _page_number_footer(document: DocxDocument) -> None:
    paragraph = document.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Content Processor guide · page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    field = paragraph.add_run()
    field.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), tag)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        field._r.append(element)  # noqa: SLF001


def _styles(document: DocxDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY
        style.font.bold = True
    section = document.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.2))


def build(source: Path = SOURCE, output: Path = OUTPUT) -> Path:
    """Convert the Markdown guide to a styled Word document."""
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    _styles(document)
    _page_number_footer(document)
    i = 0
    after_title = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            _table(document, rows)
            document.add_paragraph()
            continue
        if line.startswith("# "):
            document.add_paragraph(line[2:], style="Title")
            after_title = True
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            _add_runs(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            _add_runs(document.add_paragraph(style="List Number"), line.split(". ", 1)[1])
        else:
            paragraph = document.add_paragraph()
            _add_runs(paragraph, line)
            if after_title:  # the "Last updated" line under the title
                for run in paragraph.runs:
                    run.italic = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = GREY
                after_title = False
        i += 1
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # python-docx's template omits it; schema requires it
    document.core_properties.title = "Content Processor: A Plain-Language Guide"
    document.core_properties.author = "Rankuno"
    document.save(str(output))
    return output


if __name__ == "__main__":
    path = build()
    sys.stdout.write(f"wrote {path}\n")
